"""JLCPCB assembly files: BOM + CPL (pick & place) as Excel, from an Eagle 7.x .brd.

Per JLCPCB's help pages (checked 2026-10-08):
  BOM columns: Comment, Designator, Footprint, LCSC Part #        (.csv/.xls/.xlsx accepted)
  CPL columns: Designator, Mid X, Mid Y, Layer, Rotation          (mm; rotation in degrees, + = counter-clockwise)
  The order page has ONE BOM slot and ONE CPL slot and every CPL row carries its Layer (Top/Bottom),
  so 'Both' writes one combined pair of files. Top/Bottom write only that side's parts.

Placement comes from the .brd (the schematic has no positions). Mid X/Y is the centre of the
part's pads (not the package origin, which on some library parts sits on pin 1), transformed by the
element's rotation/mirror.

NOT verified against a real JLC order: the rotation of bottom-side parts. JLC's assembly viewer shows
every part before you pay; fix any that look turned with rotation offsets (jlc_rotations.json).
"""
from __future__ import annotations

import json
import math
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

LCSC_KEYS = ("LCSC", "LCSC_PART", "LCSC#", "JLC", "JLCPCB", "JLC_PART", "LCSC_PART_NUMBER")
DNP_RE = re.compile(r"^(DNP|DNF|NC|NOPOP|DNI)$", re.I)


@dataclass
class Placement:
    ref: str
    value: str
    package: str
    library: str
    x: float           # mm, centre of pads
    y: float
    angle: float
    bottom: bool
    attrs: Dict[str, str] = field(default_factory=dict)
    origin_offset: float = 0.0   # distance between package origin and pad centre (mm)
    has_pads: bool = True


def parse_rot(rot: str) -> Tuple[float, bool]:
    rot = rot or "R0"
    mirror = "M" in rot
    m = re.search(r"R(-?[\d.]+)", rot)
    return (float(m.group(1)) % 360 if m else 0.0), mirror


def _pad_centre(pkg: Optional[ET.Element]) -> Tuple[float, float, bool]:
    if pkg is None:
        return 0.0, 0.0, False
    xs, ys = [], []
    for s in pkg.findall("smd"):
        x, y = float(s.get("x", 0)), float(s.get("y", 0))
        dx, dy = float(s.get("dx", 0)) / 2, float(s.get("dy", 0)) / 2
        if re.search(r"R(90|270)", s.get("rot", "") or ""):
            dx, dy = dy, dx
        xs += [x - dx, x + dx]
        ys += [y - dy, y + dy]
    for p in pkg.findall("pad"):
        x, y = float(p.get("x", 0)), float(p.get("y", 0))
        r = float(p.get("diameter", 0) or float(p.get("drill", 0.8)) * 1.6) / 2
        xs += [x - r, x + r]
        ys += [y - r, y + r]
    if not xs:
        return 0.0, 0.0, False
    return (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2, True


def read_brd(path: Path) -> List[Placement]:
    root = ET.parse(path).getroot()
    board = root.find("drawing/board")
    if board is None:
        raise ValueError(f"{Path(path).name} is not an Eagle board file")
    pkgs: Dict[Tuple[str, str], ET.Element] = {}
    for lib in board.findall("libraries/library"):
        for pk in lib.findall("packages/package"):
            pkgs[(lib.get("name", ""), pk.get("name", ""))] = pk
    out = []
    for e in board.findall("elements/element"):
        angle, mirror = parse_rot(e.get("rot", "R0"))
        cx, cy, has = _pad_centre(pkgs.get((e.get("library", ""), e.get("package", ""))))
        if mirror:
            cx = -cx
        a = math.radians(angle)
        rx, ry = cx * math.cos(a) - cy * math.sin(a), cx * math.sin(a) + cy * math.cos(a)
        attrs = {a_.get("name", "").upper(): a_.get("value", "") for a_ in e.findall("attribute")}
        out.append(Placement(e.get("name", ""), e.get("value", ""), e.get("package", ""), e.get("library", ""),
                             round(float(e.get("x", 0)) + rx, 4), round(float(e.get("y", 0)) + ry, 4),
                             angle, mirror, attrs, round(math.hypot(rx, ry), 3), has))
    return out


def load_rotation_offsets(path: Path) -> Dict[str, float]:
    """{"SOT-23": 180, "U3": 90} - keys are footprint names or refs; values are added degrees."""
    try:
        return {k.upper(): float(v) for k, v in json.loads(Path(path).read_text(encoding="utf-8")).items()}
    except Exception:
        return {}


def build(design, placements: List[Placement], side: str = "top", bottom_mode: str = "eagle",
          offsets: Optional[Dict[str, float]] = None) -> Dict:
    """side: top | bottom | both. bottom_mode: 'eagle' (angle as shown in Eagle) or 'mirror' (360-angle)."""
    offsets = offsets or {}
    want = {"top": (False,), "bottom": (True,), "both": (False, True)}[side]
    parts = getattr(design, "parts", {}) or {}
    cpl, groups, warn = [], {}, []
    skipped_dnp, no_pads, missing_lcsc = [], [], []
    for pl in sorted(placements, key=lambda p: (re.sub(r"\d", "", p.ref), int(re.sub(r"\D", "", p.ref) or 0))):
        if pl.bottom not in want:
            continue
        spec = parts.get(pl.ref.upper())
        value = (spec.value if spec and spec.value else pl.value) or (spec.deviceset if spec else "")
        if not pl.has_pads:
            no_pads.append(pl.ref)
            continue
        if (spec and spec.dnp) or DNP_RE.match(value or "") or pl.attrs.get("DNP") or pl.attrs.get("POPULATE", "").lower() in ("no", "0", "false"):
            skipped_dnp.append(pl.ref)
            continue
        lcsc = (spec.lcsc if spec and spec.lcsc else "") or next((pl.attrs[k] for k in LCSC_KEYS if pl.attrs.get(k)), "")
        if not lcsc:
            missing_lcsc.append(pl.ref)
        rot = pl.angle
        if pl.bottom and bottom_mode == "mirror":
            rot = (360 - rot) % 360
        rot = (rot + offsets.get(pl.ref.upper(), offsets.get(pl.package.upper(), 0))) % 360
        cpl.append({"Designator": pl.ref, "Mid X": round(pl.x, 4), "Mid Y": round(pl.y, 4),
                    "Layer": "Bottom" if pl.bottom else "Top", "Rotation": round(rot, 2)})
        if pl.origin_offset > 0.05:
            warn.append(f"{pl.ref}: package origin is {pl.origin_offset} mm from the pad centre - centre used")
        groups.setdefault((value, pl.package, lcsc), []).append(pl.ref)
    bom = [{"Comment": k[0], "Designator": ",".join(v), "Footprint": k[1], "LCSC Part #": k[2], "Qty": len(v)}
           for k, v in groups.items()]
    design_refs = set(parts)
    board_refs = {p.ref.upper() for p in placements}
    not_on_board = sorted(r for r in design_refs - board_refs)
    return {"bom": bom, "cpl": cpl, "missing_lcsc": missing_lcsc, "dnp": skipped_dnp, "no_pads": no_pads,
            "not_on_board": not_on_board, "warnings": warn,
            "sides_present": sorted({("Bottom" if p.bottom else "Top") for p in placements if p.has_pads})}


def write_xlsx(rows: List[Dict], columns: List[str], path: Path, title: str, highlight_empty: Optional[str] = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FFC7CE")
    for r in rows:
        ws.append([r.get(c, "") for c in columns])
        if highlight_empty and not r.get(highlight_empty):
            for c in ws[ws.max_row]:
                c.fill = red
    for i, col in enumerate(columns, 1):
        width = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    ws.freeze_panes = "A2"
    wb.save(path)


def export_paths(brd: Path, out_dir: Path, side: str) -> Tuple[Path, Path]:
    stem = f"{Path(brd).stem}_{side}"
    return Path(out_dir) / f"BOM_{stem}.xlsx", Path(out_dir) / f"CPL_{stem}.xlsx"


def export(design, brd: Path, out_dir: Path, side: str, bottom_mode: str = "eagle",
           offsets_path: Optional[Path] = None, bom_path: Optional[Path] = None, cpl_path: Optional[Path] = None) -> Dict:
    res = build(design, read_brd(brd), side, bottom_mode, load_rotation_offsets(offsets_path) if offsets_path else {})
    out_dir.mkdir(parents=True, exist_ok=True)
    bom_p, cpl_p = export_paths(brd, out_dir, side)
    bom_p, cpl_p = Path(bom_path or bom_p), Path(cpl_path or cpl_p)
    write_xlsx(res["bom"], ["Comment", "Designator", "Footprint", "LCSC Part #", "Qty"], bom_p, "BOM",
               highlight_empty="LCSC Part #")
    write_xlsx(res["cpl"], ["Designator", "Mid X", "Mid Y", "Layer", "Rotation"], cpl_p, "CPL")
    res.update(bom_path=str(bom_p), cpl_path=str(cpl_p))
    return res
