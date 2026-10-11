import os, sys, tempfile, json, zipfile, shutil, glob
os.environ["QT_QPA_PLATFORM"] = "offscreen"
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1])); sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_engine as te
from fakes import FakeProvider
from eagleai.lbr_builder import build_lbr
from eagleai.tools import Project, Settings, dispatch
from eagleai.agent import Agent
from eagleai import jlc_export, ul_import
import xml.etree.ElementTree as ET

tmp = Path(tempfile.mkdtemp())
lib = tmp / "lbr"
for s in (te.BUCK, te.SOT, te.RES): build_lbr(s, lib)
st = Settings(); st.lib_dirs = [str(lib)]; st.converted_dir = str(tmp / "conv"); st.memory_path = str(tmp / "mem.json")
st.scripts_dir = str(tmp / "scripts")
p = Project(tmp / "proj", st); p.load_libraries()
ag = Agent(p)

# 1) Claude builds, Grok continues: Grok sees a COMPACT history incl. Claude's tool calls
c = FakeProvider([("Placing.", [("add_parts", {"parts": [{"library": "LDO5", "deviceset": "LDO5"}]}),
                                ("add_passive", {"kind": "R", "value": "10k", "package": "0603", "count": 2})]),
                  ("Placed U1, R1, R2.", [])], brand="claude", model="claude-sonnet-5-5")
ag.run("place an LDO and two resistors", provider=c, brand="claude")
g = FakeProvider([("Connecting.", [("connect_nets", {"nets": {"VIN": ["U1.IN", "R1.1"], "GND": ["U1.GND", "R2.2"]}})]),
                  ("Done.", [])], brand="xai", model="grok-4.7")
ag.run("now wire VIN and GND", provider=g, brand="xai")
hist = g.seen_history
print(json.dumps(hist, indent=1)[:900])
assert hist[0]["role"] == "user" and "add_parts" in hist[1]["text"] and "[Claude" in hist[1]["text"]
assert p.design.nets["VIN"] == ["U1.IN", "R1.1"]
assert len(json.dumps(hist)) < 1500, len(json.dumps(hist))
# persisted: a brand-new Agent (app restart) sees the same log
ag2 = Agent(p); assert len(ag2.visible_log()) == len(ag.visible_log()) > 4
assert p.costs.price_for("grok-4.7", 1)["input"] == 2.0

# 2) review is read-only even if the reviewer tries to edit
r = FakeProvider([("Checking.", [("get_design", {}), ("remove_part", {"ref": "U1"})]), ("PROBLEMS: none", [])],
                 brand="openai", model="gpt-5.6-terra")
res = ag.run("Double-check", provider=r, brand="openai", review=True)
assert "U1" in p.design.parts and "remove_part" not in r.seen_tools, r.seen_tools
assert ag.log[-1]["t"] == "review"
print("review tools:", r.seen_tools)

# 3) new chat: compact history restarts, file keeps everything
n_before = len(ag.log); ag.new_chat()
assert ag.compact_history() == [] and len(Agent(p).log) == n_before + 1

# 4) scripts: write -> refused without approval -> approved -> runs and edits design
code = """import json, os
d = json.load(open(os.environ['EAGLEAI_DESIGN']))
print('parts:', len(d['parts']))
d['notes'] = 'touched by script'
json.dump(d, open(os.environ['EAGLEAI_DESIGN'], 'w'))
"""
print(dispatch(p, "write_script", {"name": "count parts", "code": code, "description": "Count parts"}))
p.ask_script = lambda n, c: False
print(dispatch(p, "run_script", {"name": "count_parts"}))
asked = []
p.ask_script = lambda n, c: asked.append(n) or True
out = dispatch(p, "run_script", {"name": "count_parts"}); print(out)
assert out["stdout"].strip() == "parts: 3" and p.design.notes == "touched by script"
out = dispatch(p, "run_script", {"name": "count_parts"}); assert len(asked) == 1     # approval remembered
print(dispatch(p, "list_scripts", {}))

# 5) JLC export from a synthetic .brd (top + bottom, origin at pin 1 on one part)
brd = tmp / "proj" / "board.brd"
root = ET.Element("eagle", version="7.4.0"); bd = ET.SubElement(ET.SubElement(root, "drawing"), "board")
libs = ET.SubElement(bd, "libraries")
L = ET.SubElement(libs, "library", name="LDO5"); pk = ET.SubElement(ET.SubElement(L, "packages"), "package", name="SOT23-5")
for i, (x, y) in enumerate([(-0.95, -1.3), (0, -1.3), (0.95, -1.3), (0.95, 1.3), (-0.95, 1.3)]):
    ET.SubElement(pk, "smd", name=str(i + 1), x=str(x), y=str(y), dx="0.6", dy="1.1", layer="1")
L2 = ET.SubElement(libs, "library", name="R"); pk2 = ET.SubElement(ET.SubElement(L2, "packages"), "package", name="R0603")
ET.SubElement(pk2, "smd", name="1", x="0", y="0", dx="0.8", dy="0.9", layer="1")      # origin on pad 1
ET.SubElement(pk2, "smd", name="2", x="1.6", y="0", dx="0.8", dy="0.9", layer="1")
els = ET.SubElement(bd, "elements")
ET.SubElement(els, "element", name="U1", library="LDO5", package="SOT23-5", value="LDO5", x="10", y="10", rot="R90")
ET.SubElement(els, "element", name="R1", library="R", package="R0603", value="10k", x="20", y="5", rot="R0")
ET.SubElement(els, "element", name="R2", library="R", package="R0603", value="10k", x="20", y="15", rot="MR180")
ET.ElementTree(root).write(brd)
p.design.parts["R1"].lcsc = "C25804"
for side in ("top", "bottom", "both"):
    res = jlc_export.export(p.design, brd, tmp / "proj" / "jlc", side)
    print(side, res["cpl"], res["bom"], res["missing_lcsc"])
r1 = next(r for r in res["cpl"] if r["Designator"] == "R1"); assert (r1["Mid X"], r1["Mid Y"]) == (20.8, 5.0), r1
r2 = next(r for r in res["cpl"] if r["Designator"] == "R2"); assert r2["Layer"] == "Bottom" and r2["Mid X"] == 20.8, r2
assert {r["Layer"] for r in res["cpl"]} == {"Top", "Bottom"}
from openpyxl import load_workbook
ws = load_workbook(res["bom_path"]).active
print([c.value for c in ws[1]], [c.value for c in ws[2]])
assert [c.value for c in ws[1]][:4] == ["Comment", "Designator", "Footprint", "LCSC Part #"]
ws = load_workbook(res["cpl_path"]).active; assert [c.value for c in ws[1]] == ["Designator", "Mid X", "Mid Y", "Layer", "Rotation"]

# 6) Ultra Librarian zips: the real sample + Fusion-only and script-only variants built from it
sample = os.environ.get("UL_SAMPLE", "/tmp/claude-0/ulzip/in.zip")
if Path(sample).exists():
    src = zipfile.ZipFile(sample)
    def variant(name, keep):
        zp = tmp / name
        with zipfile.ZipFile(zp, "w") as z:
            for n in src.namelist():
                if keep(n): z.writestr(n, src.read(n))
        return zp
    for zp in (Path(sample), variant("ul_fusion_only.zip", lambda n: "Eagle/" not in n),
               variant("ul_script_only.zip", lambda n: "Fusion360/" not in n)):
        dest = tmp / "ullib" / zp.stem
        r = ul_import.import_zip(zp, dest)
        print(zp.name, "->", r["part"], r["report"]); assert r["part"] == "LM7805CT"
        dv = ET.parse(r["library"]).getroot()
        assert dv.get("version") == "7.4.0"
        assert sorted((c.get("pin"), c.get("pad")) for c in dv.iter("connect")) == [("GND", "2"), ("INPUT", "1"), ("OUTPUT", "3")]
    from eagleai.library import LibraryIndex
    ix = LibraryIndex([tmp / "ullib"]); ix.scan()
    print(ix.stats()); assert len(ix.duplicates) == 2 and not any("fusion360" in str(f) for f in ix._files.values())
    # the .sch writer must not embed Fusion-only stuff into an Eagle 7.4 schematic
    pz = Project(tmp / "projz", Settings(lib_dirs=[str(tmp / "ullib")], converted_dir=str(tmp / "c2"),
                                          memory_path=str(tmp / "m2.json"), scripts_dir=str(tmp / "s2")))
    pz.load_libraries()
    pz.design.add_part(pz.index, "LM7805CT", "LM7805CT")
    pz.design.connect(pz.index, "VIN", ["U1.INPUT"]); pz.design.connect(pz.index, "GND", ["U1.GND"])
    print(dispatch(pz, "generate_sch", {"filename": "reg"}))
else:
    print("UL sample not present - skipped")

# 7) UI smoke: recent list, brand switch, chat rendering
from PySide6.QtWidgets import QApplication, QMessageBox
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.No)
QMessageBox.information = staticmethod(lambda *a, **k: None)
from PySide6.QtCore import QSettings
QSettings("EagleAI", "IDE").clear()
from eagleai.app import MainWindow
import time
app = QApplication([])
w = MainWindow(); w.settings.lib_dirs = st.lib_dirs; w.settings.converted_dir = st.converted_dir
w.settings.scripts_dir = st.scripts_dir; w.settings.memory_path = st.memory_path
w.open_project(tmp / "proj")
for _ in range(80): app.processEvents(); time.sleep(0.02)
w.open_project(tmp / "projz" if (tmp / "projz").exists() else tmp / "proj2")
for _ in range(40): app.processEvents(); time.sleep(0.02)
w.open_project(tmp / "proj")
for _ in range(40): app.processEvents(); time.sleep(0.02)
print("recent:", w.recent()); assert len(w.recent()) == 2 and w.recent()[0].endswith("proj")
w.brand_box.setCurrentIndex(2); assert w.settings.brand == "xai" and w.model_box.currentText() == "grok-4.7"
w.brand_box.setCurrentIndex(0)
w._fill_project_menu(); print([a.text() for a in w.proj_menu.actions()][:4])
w.show(); app.processEvents(); w.grab().save(str(tmp / "v04.png")); print("shot", tmp / "v04.png")
print("ALL OK")
