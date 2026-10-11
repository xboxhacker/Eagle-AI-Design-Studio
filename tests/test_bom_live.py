import os, sys, tempfile, time
os.environ["QT_QPA_PLATFORM"] = "offscreen"
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1])); sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_engine as te
import xml.etree.ElementTree as ET
from eagleai.lbr_builder import build_lbr
from eagleai.tools import Project, Settings, dispatch
from eagleai import bom_sync

tmp = Path(tempfile.mkdtemp())
for s in (te.SOT, te.RES): build_lbr(s, tmp / "lbr")
st = Settings(lib_dirs=[str(tmp / "lbr")], converted_dir=str(tmp / "c"), memory_path=str(tmp / "m.json"), scripts_dir=str(tmp / "s"))
p = Project(tmp / "proj", st); p.load_libraries()
dispatch(p, "add_parts", {"parts": [{"library": "LDO5", "deviceset": "LDO5"}]})
dispatch(p, "add_passive", {"kind": "R", "value": "10k", "package": "0603", "count": 3})
# board with all 4 parts on top
brd = tmp / "proj" / "b.brd"
root = ET.Element("eagle", version="7.4.0"); bd = ET.SubElement(ET.SubElement(root, "drawing"), "board")
libs = ET.SubElement(bd, "libraries")
for ln, pn in (("LDO5", "SOT23-5"), ("R", "R0603")):
    pk = ET.SubElement(ET.SubElement(ET.SubElement(libs, "library", name=ln), "packages"), "package", name=pn)
    ET.SubElement(pk, "smd", name="1", x="-0.8", y="0", dx="0.8", dy="0.9", layer="1"); ET.SubElement(pk, "smd", name="2", x="0.8", y="0", dx="0.8", dy="0.9", layer="1")
els = ET.SubElement(bd, "elements")
for i, (n, l, pk) in enumerate((("U1", "LDO5", "SOT23-5"), ("R1", "R", "R0603"), ("R2", "R", "R0603"), ("R3", "R", "R0603"))):
    ET.SubElement(els, "element", name=n, library=l, package=pk, value="", x=str(5 + 5 * i), y="5", rot="R0")
ET.ElementTree(root).write(brd)

# 1) export, then status = up to date
for r in ("R1", "R2", "R3"): p.design.parts[r].lcsc = "C25804"
res = bom_sync.export_jlc(p, brd, "both")
st0 = bom_sync.export_status(p)[0]; print(st0["name"], st0["up_to_date"], st0["differences"]); assert st0["up_to_date"]
# 2) change the design -> export is out of date, with exact differences
dispatch(p, "set_value", {"ref": "R2", "value": "4.7k"}); dispatch(p, "set_value", {"ref": "R3", "dnp": True})
st1 = bom_sync.export_status(p)[0]; print(st1["differences"]); assert not st1["up_to_date"] and len(st1["differences"]) == 2
# 3) re-export (user picks Overwrite) -> fine again; then edit the xlsx "in Excel" -> edited_outside + pull
p.guard.ask = lambda *a: ("overwrite", False)
bom_sync.export_jlc(p, brd, "both")
assert bom_sync.export_status(p)[0]["up_to_date"]
from openpyxl import load_workbook
time.sleep(1.1)
wb = load_workbook(res["bom_path"]); ws = wb.active
for row in ws.iter_rows(min_row=2):
    if row[1].value == "U1": row[3].value = "C6186"
wb.save(res["bom_path"])
st2 = bom_sync.export_status(p)[0]; print(st2["edited_outside"], st2["differences"]); assert st2["edited_outside"]
ch = bom_sync.pull_candidates(p, Path(res["bom_path"])); print(ch); assert ch == [("U1", "lcsc", "", "C6186")]
bom_sync.apply_pull(p, ch); bom_sync.mark_seen(p, Path(res["bom_path"]))
st3 = bom_sync.export_status(p)[0]; print(st3); assert st3["up_to_date"] and not st3["edited_outside"]
assert p.design.parts["U1"].lcsc == "C6186"

# 4) the widget: live rows, edit -> design, stale price hidden, locked while busy
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QMessageBox; QMessageBox.warning = staticmethod(lambda *a, **k: None)
from eagleai.bom_view import BomEditor, COLS
app = QApplication([])
p.save_state(prices={"R1": {"lcsc": "C25804", "unit": 0.002, "stock": 5000, "ts": time.time()},
                     "U1": {"lcsc": "C9999", "unit": 0.5, "stock": 10, "ts": time.time()}})   # U1 price is for an old LCSC
edits, busy = [], [False]
ed = BomEditor(lambda: p, edits.append, lambda: busy[0], {})
ed._do_refresh()
def row_of(ref):
    for i in range(ed.table.rowCount()):
        if ref in ed.table.item(i, 0).text().split(", "): return i
cells = lambda ref: [ed.table.item(row_of(ref), j).text() for j in range(len(COLS))]
print(cells("U1")); print(cells("R1"))
assert cells("U1")[COLS.index("Unit $")] == "" and cells("R1")[COLS.index("Unit $")] == "0.0020"
ed.table.item(row_of("R1"), COLS.index("MPN")).setText("0603WAF1002T5E")
print(edits, p.design.parts["R1"].mpn); assert p.design.parts["R1"].mpn == "0603WAF1002T5E" and edits
ed.table.item(row_of("U1"), COLS.index("LCSC")).setText("hello")       # rejected (message box suppressed)
assert p.design.parts["U1"].lcsc == "C6186"
busy[0] = True; ed._do_refresh()
assert not (ed.table.item(row_of("R1"), COLS.index("Value")).flags() & Qt.ItemIsEditable)
print("banner:", ed.banner.text()[:160]); print(ed.total.text())
print("ALL OK")
